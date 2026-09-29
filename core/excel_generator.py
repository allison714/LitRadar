"""
Excel Knowledge Base Generator.
LitRadar Suite.
"""

from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import EXCEL_PATH


class ExcelKnowledgeBase:
    def __init__(self, excel_path: Path = EXCEL_PATH):
        self.excel_path = excel_path
        self.header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        self.link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        self.regular_font = Font(name="Calibri", size=10)
        self.thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

    def _apply_header_style(self, row):
        for cell in row:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _auto_fit_columns(self, ws, max_width_cap=55):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_width_cap)

    def update_knowledge_base(self, new_papers: List[Dict[str, Any]]):
        if self.excel_path.exists():
            wb = openpyxl.load_workbook(self.excel_path)
        else:
            wb = openpyxl.Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

        # 1. Weekly_Arrivals
        if "Weekly_Arrivals" in wb.sheetnames:
            ws_weekly = wb["Weekly_Arrivals"]
            ws_weekly.delete_rows(1, ws_weekly.max_row + 1)
        else:
            ws_weekly = wb.create_sheet(title="Weekly_Arrivals", index=0)

        weekly_headers = [
            "Track", "Date Added", "Year", "Title", "Local PDF", "Journal", "Authors & Lab",
            "Citations", "TL;DR Finding", "Key Targets", "Methods", "Zotero Tags"
        ]
        ws_weekly.append(weekly_headers)
        self._apply_header_style(ws_weekly[1])
        ws_weekly.row_dimensions[1].height = 26

        today_str = datetime.now().strftime("%Y-%m-%d")
        for idx, paper in enumerate(new_papers, start=2):
            pdf_path_val = paper.get("local_pdf_path", "")
            pdf_display = "📄 Open PDF" if pdf_path_val else "Online Only"
            row_data = [
                paper.get("track_name", "General"),
                today_str,
                paper.get("year", datetime.now().year),
                paper.get("title", ""),
                pdf_display,
                paper.get("journal", ""),
                paper.get("authors", ""),
                paper.get("citations", 0),
                paper.get("tldr", ""),
                ", ".join(paper.get("targets", [])),
                ", ".join(paper.get("methods", [])),
                ", ".join(paper.get("tags", [])),
            ]
            ws_weekly.append(row_data)

            row_cells = ws_weekly[idx]
            is_even = (idx % 2 == 0)
            for col_idx, cell in enumerate(row_cells, start=1):
                cell.font = self.regular_font
                cell.border = self.thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if is_even:
                    cell.fill = self.zebra_fill
                if col_idx == 4 and paper.get("url"):
                    cell.hyperlink = paper["url"]
                    cell.font = self.link_font
                elif col_idx == 5 and pdf_path_val:
                    cell.hyperlink = f"file:///{str(pdf_path_val).replace(chr(92), '/')}"
                    cell.font = self.link_font

        self._auto_fit_columns(ws_weekly)

        # 2. Master_Archive (Year & Citation Sorted)
        archive_headers = [
            "Date Added", "DOI / ID", "Year", "Citations", "Title", "Journal", "Track",
            "Key Targets", "Methods", "TL;DR Summary", "Zotero Tags"
        ]

        existing_entries = {}
        if "Master_Archive" in wb.sheetnames:
            ws_archive = wb["Master_Archive"]
            header_row = [str(c).strip().lower() for c in (next(ws_archive.iter_rows(min_row=1, max_row=1, values_only=True), None) or [])]
            has_date_col = len(header_row) > 0 and "date" in header_row[0]

            for row in ws_archive.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue
                if has_date_col:
                    date_val = str(row[0] or "") if len(row) > 0 else today_str
                    doi_val = row[1] if len(row) > 1 else ""
                    year_val = row[2] if len(row) > 2 else ""
                    cites_val = row[3] if len(row) > 3 else 0
                    title_val = row[4] if len(row) > 4 else ""
                    journal_val = row[5] if len(row) > 5 else ""
                    track_val = row[6] if len(row) > 6 else ""
                    targets_val = row[7].split(", ") if len(row) > 7 and row[7] else []
                    methods_val = row[8].split(", ") if len(row) > 8 and row[8] else []
                    tldr_val = row[9] if len(row) > 9 else ""
                    tags_val = row[10].split(", ") if len(row) > 10 and row[10] else []
                else:
                    date_val = "Historical Archive"
                    doi_val = row[0] if len(row) > 0 else ""
                    year_val = row[1] if len(row) > 1 else ""
                    cites_val = row[2] if len(row) > 2 else 0
                    title_val = row[3] if len(row) > 3 else ""
                    journal_val = row[4] if len(row) > 4 else ""
                    track_val = row[5] if len(row) > 5 else ""
                    targets_val = row[6].split(", ") if len(row) > 6 and row[6] else []
                    methods_val = row[7].split(", ") if len(row) > 7 and row[7] else []
                    tldr_val = row[8] if len(row) > 8 else ""
                    tags_val = row[9].split(", ") if len(row) > 9 and row[9] else []

                if doi_val or title_val:
                    key = str(doi_val or title_val).lower().strip()
                    existing_entries[key] = {
                        "date_added": date_val,
                        "doi": doi_val,
                        "year": year_val,
                        "citations": cites_val,
                        "title": title_val,
                        "journal": journal_val,
                        "track_name": track_val,
                        "targets": targets_val,
                        "methods": methods_val,
                        "tldr": tldr_val,
                        "tags": tags_val,
                        "url": f"https://doi.org/{doi_val}" if "10." in str(doi_val) else "",
                    }
            wb.remove(ws_archive)

        for p in new_papers:
            ident = p.get("doi") or p.get("pmid") or p.get("title")
            if ident:
                key = str(ident).lower().strip()
                if key in existing_entries:
                    # Keep original date_added if already present
                    p["date_added"] = existing_entries[key].get("date_added", today_str)
                else:
                    p["date_added"] = p.get("date_added", today_str)
                existing_entries[key] = p

        ws_archive = wb.create_sheet(title="Master_Archive")
        ws_archive.append(archive_headers)
        self._apply_header_style(ws_archive[1])
        ws_archive.row_dimensions[1].height = 26

        sorted_archive = sorted(
            existing_entries.values(),
            key=lambda x: (x.get("year", 0) or 0, x.get("citations", 0) or 0),
            reverse=True
        )

        for idx, paper in enumerate(sorted_archive, start=2):
            doi_val = paper.get("doi") or paper.get("pmid") or ""
            row_data = [
                paper.get("date_added", today_str),
                doi_val,
                paper.get("year", ""),
                paper.get("citations", 0),
                paper.get("title", ""),
                paper.get("journal", ""),
                paper.get("track_name", ""),
                ", ".join(paper.get("targets", [])) if isinstance(paper.get("targets"), list) else str(paper.get("targets") or ""),
                ", ".join(paper.get("methods", [])) if isinstance(paper.get("methods"), list) else str(paper.get("methods") or ""),
                paper.get("tldr", ""),
                ", ".join(paper.get("tags", [])) if isinstance(paper.get("tags"), list) else str(paper.get("tags") or ""),
            ]
            ws_archive.append(row_data)

            row_cells = ws_archive[idx]
            is_even = (idx % 2 == 0)
            for col_idx, cell in enumerate(row_cells, start=1):
                cell.font = self.regular_font
                cell.border = self.thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if is_even:
                    cell.fill = self.zebra_fill
                if col_idx == 5 and paper.get("url"):
                    cell.hyperlink = paper["url"]
                    cell.font = self.link_font

        self._auto_fit_columns(ws_archive)

        # 3. PanExM_Antibodies
        ab_headers = [
            "Target Antigen", "Host Species", "Vendor", "Catalog No.",
            "Dilution", "Pan-ExM Staining Notes", "Reference Paper"
        ]

        if "PanExM_Antibodies" in wb.sheetnames:
            ws_ab = wb["PanExM_Antibodies"]
        else:
            ws_ab = wb.create_sheet(title="PanExM_Antibodies")
            ws_ab.append(ab_headers)
            self._apply_header_style(ws_ab[1])
            ws_ab.row_dimensions[1].height = 26

        current_ab_count = ws_ab.max_row
        for paper in new_papers:
            for ab in paper.get("pan_exm_antibodies", []):
                current_ab_count += 1
                row_data = [
                    ab.get("target", "Unknown"),
                    ab.get("host", "Unknown"),
                    ab.get("vendor", "Unspecified"),
                    ab.get("catalog_no", "N/A"),
                    ab.get("dilution", "N/A"),
                    ab.get("notes", "Extracted from literature"),
                    paper.get("title", ""),
                ]
                ws_ab.append(row_data)
                row_cells = ws_ab[current_ab_count]
                for cell in row_cells:
                    cell.font = self.regular_font
                    cell.border = self.thin_border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

        self._auto_fit_columns(ws_ab)

        wb.save(self.excel_path)
        print(f"[Excel Knowledge Base] Updated: {self.excel_path}")
